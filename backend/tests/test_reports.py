"""Attendance reporting + export (issue #2 "Pelaporan & ekspor").

Seeds attendance records directly under a tenant context (RLS-bound session via
``_set_tenant``), then exercises the recap JSON endpoints and the CSV / XLSX
export through the ASGI client. Cross-tenant isolation is asserted last: a
tenant must only ever see its own records.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.db.session import SessionFactory, _set_tenant
from app.models import AttendanceRecord, AttendanceStatus, AttendanceType, Role, Tenant, User


async def _token(client: AsyncClient, **payload) -> str:
    resp = await client.post("/api/v1/auth/login", json=payload)
    assert resp.status_code == 200, resp.text
    return resp.json()["data"]["access_token"]


async def _onboard(client: AsyncClient, owner_hdr: dict, slug: str) -> dict:
    resp = await client.post(
        "/api/v1/tenants",
        headers=owner_hdr,
        json={
            "name": f"Sekolah {slug}",
            "slug": slug,
            "admin_username": f"admin-{slug}",
            "admin_password": "adminpass123",
            "admin_full_name": "Admin",
        },
    )
    assert resp.status_code == 201, resp.text
    admin = await _token(
        client, username=f"admin-{slug}", password="adminpass123", tenant_slug=slug
    )
    return {"Authorization": f"Bearer {admin}"}


async def _tenant_id(slug: str) -> str:
    from sqlalchemy import select

    async with SessionFactory() as s:
        await _set_tenant(s, None)
        return (await s.execute(select(Tenant.id).where(Tenant.slug == slug))).scalar_one()


async def _seed_user(tenant_id: str, full_name: str, username: str) -> str:
    async with SessionFactory() as s:
        await _set_tenant(s, tenant_id)
        user = User(
            tenant_id=tenant_id,
            full_name=full_name,
            username=username,
            role=Role.end_user,
            is_active=True,
        )
        s.add(user)
        await s.commit()
        await s.refresh(user)
        return user.id


async def _seed_record(
    tenant_id: str,
    user_id: str,
    *,
    att_type: AttendanceType,
    status: AttendanceStatus,
    occurred_at: datetime,
    liveness_score: float | None = None,
    device_id: str | None = None,
) -> None:
    async with SessionFactory() as s:
        await _set_tenant(s, tenant_id)
        s.add(
            AttendanceRecord(
                tenant_id=tenant_id,
                user_id=user_id,
                type=att_type,
                status=status,
                occurred_at=occurred_at,
                liveness_score=liveness_score,
                device_id=device_id,
            )
        )
        await s.commit()


@pytest.mark.asyncio
async def test_daily_and_monthly_recap(client: AsyncClient, super_admin):
    owner = await _token(client, username="owner", password="ownerpass123")
    hdr = await _onboard(client, {"Authorization": f"Bearer {owner}"}, "rep-a")
    tid = await _tenant_id("rep-a")
    alice = await _seed_user(tid, "Alice", "alice")
    bob = await _seed_user(tid, "Bob", "bob")

    # Day under test: 2026-06-10.
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.late,
        occurred_at=datetime(2026, 6, 10, 9, 0, tzinfo=UTC),
        liveness_score=0.95,
    )
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_out,
        status=AttendanceStatus.early_leave,
        occurred_at=datetime(2026, 6, 10, 15, 0, tzinfo=UTC),
    )
    await _seed_record(
        tid,
        bob,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 6, 10, 8, 0, tzinfo=UTC),
    )
    # Different day, same month.
    await _seed_record(
        tid,
        bob,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 6, 20, 8, 0, tzinfo=UTC),
    )
    # Different month -> excluded from June.
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 7, 1, 8, 0, tzinfo=UTC),
    )

    # --- Daily recap for 2026-06-10 ---
    resp = await client.get(
        "/api/v1/reports/attendance/daily", headers=hdr, params={"date": "2026-06-10"}
    )
    assert resp.status_code == 200, resp.text
    daily = resp.json()["data"]
    assert daily["total_records"] == 3
    assert daily["late"] == 1
    assert daily["early_leave"] == 1
    assert daily["on_time"] == 1
    assert daily["check_in"] == 2
    assert daily["check_out"] == 1
    # Per-user (sorted by name): Alice then Bob.
    assert [u["full_name"] for u in daily["users"]] == ["Alice", "Bob"]
    alice_recap = daily["users"][0]
    assert alice_recap["late"] == 1 and alice_recap["early_leave"] == 1
    assert alice_recap["check_in"] == 1 and alice_recap["check_out"] == 1

    # --- Monthly recap for 2026-06 (4 records, July excluded) ---
    resp = await client.get(
        "/api/v1/reports/attendance/monthly", headers=hdr, params={"month": "2026-06"}
    )
    assert resp.status_code == 200, resp.text
    monthly = resp.json()["data"]
    assert monthly["total_records"] == 4
    assert monthly["on_time"] == 2  # bob x2
    assert monthly["late"] == 1
    assert monthly["early_leave"] == 1
    assert monthly["check_in"] == 3


@pytest.mark.asyncio
async def test_bad_date_params_return_422(client: AsyncClient, super_admin):
    owner = await _token(client, username="owner", password="ownerpass123")
    hdr = await _onboard(client, {"Authorization": f"Bearer {owner}"}, "rep-bad")

    assert (
        await client.get(
            "/api/v1/reports/attendance/daily", headers=hdr, params={"date": "13-06-2026"}
        )
    ).status_code == 422
    assert (
        await client.get(
            "/api/v1/reports/attendance/monthly", headers=hdr, params={"month": "2026-13"}
        )
    ).status_code == 422


@pytest.mark.asyncio
async def test_csv_export(client: AsyncClient, super_admin):
    owner = await _token(client, username="owner", password="ownerpass123")
    hdr = await _onboard(client, {"Authorization": f"Bearer {owner}"}, "rep-csv")
    tid = await _tenant_id("rep-csv")
    alice = await _seed_user(tid, "Alice", "alice")
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 6, 10, 8, 0, tzinfo=UTC),
        liveness_score=0.91,
        device_id=None,
    )
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_out,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 6, 10, 17, 0, tzinfo=UTC),
    )

    resp = await client.get(
        "/api/v1/reports/attendance/export",
        headers=hdr,
        params={"from": "2026-06-01", "to": "2026-06-30", "format": "csv"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]
    assert ".csv" in resp.headers["content-disposition"]

    lines = resp.text.strip().splitlines()
    assert lines[0] == "occurred_at,user_id,full_name,type,status,liveness_score,device_id"
    assert len(lines) == 3  # header + 2 records
    assert "Alice" in lines[1]
    assert "check_in" in lines[1]


@pytest.mark.asyncio
async def test_xlsx_export(client: AsyncClient, super_admin):
    owner = await _token(client, username="owner", password="ownerpass123")
    hdr = await _onboard(client, {"Authorization": f"Bearer {owner}"}, "rep-xlsx")
    tid = await _tenant_id("rep-xlsx")
    alice = await _seed_user(tid, "Alice", "alice")
    await _seed_record(
        tid,
        alice,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.late,
        occurred_at=datetime(2026, 6, 10, 9, 30, tzinfo=UTC),
        liveness_score=0.88,
    )

    resp = await client.get(
        "/api/v1/reports/attendance/export",
        headers=hdr,
        params={"from": "2026-06-01", "to": "2026-06-30", "format": "xlsx"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers["content-disposition"]
    assert ".xlsx" in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == (
        "occurred_at",
        "user_id",
        "full_name",
        "type",
        "status",
        "liveness_score",
        "device_id",
    )
    assert len(rows) == 2  # header + 1 record
    assert rows[1][2] == "Alice"
    assert rows[1][4] == "late"


@pytest.mark.asyncio
async def test_reports_are_tenant_isolated(client: AsyncClient, super_admin):
    """Tenant B must never see tenant A's attendance in any report (RLS)."""
    owner = await _token(client, username="owner", password="ownerpass123")
    ohdr = {"Authorization": f"Bearer {owner}"}

    a_hdr = await _onboard(client, ohdr, "iso-rep-a")
    a_tid = await _tenant_id("iso-rep-a")
    a_user = await _seed_user(a_tid, "Alice", "alice")
    await _seed_record(
        a_tid,
        a_user,
        att_type=AttendanceType.check_in,
        status=AttendanceStatus.on_time,
        occurred_at=datetime(2026, 6, 10, 8, 0, tzinfo=UTC),
    )

    b_hdr = await _onboard(client, ohdr, "iso-rep-b")

    # Tenant A sees its record.
    a_resp = await client.get(
        "/api/v1/reports/attendance/daily", headers=a_hdr, params={"date": "2026-06-10"}
    )
    assert a_resp.json()["data"]["total_records"] == 1

    # Tenant B sees nothing for the same day.
    b_resp = await client.get(
        "/api/v1/reports/attendance/daily", headers=b_hdr, params={"date": "2026-06-10"}
    )
    assert b_resp.status_code == 200, b_resp.text
    assert b_resp.json()["data"]["total_records"] == 0

    # B's CSV export is empty (header only).
    b_csv = await client.get(
        "/api/v1/reports/attendance/export",
        headers=b_hdr,
        params={"from": "2026-06-01", "to": "2026-06-30", "format": "csv"},
    )
    assert b_csv.status_code == 200
    assert len(b_csv.text.strip().splitlines()) == 1  # header only
